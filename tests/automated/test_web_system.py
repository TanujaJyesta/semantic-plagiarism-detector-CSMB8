"""Comprehensive automated test suite for explainable semantic plagiarism detection."""

from __future__ import annotations

import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
from docx import Document

TEST_DIRECTORY = tempfile.TemporaryDirectory()
TEST_DB_PATH = Path(TEST_DIRECTORY.name) / "test.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB_PATH}"

from backend.app import app  # noqa: E402
from backend.config import Settings  # noqa: E402
from backend.database.database import configure_database, initialize_database  # noqa: E402
from backend.services.database_service import (  # noqa: E402
    save_analysis_to_database,
    save_web_analysis_to_database,
)
from backend.services.document_processor import (  # noqa: E402
    extract_docx_text,
    extract_pdf_text,
    extract_text,
    extract_txt_text,
    normalize_text,
    process_document,
    split_sentences,
)
from backend.services.evidence_analyzer import analyze_evidence  # noqa: E402
from backend.services.exceptions import (  # noqa: E402
    ConfigurationError,
    ExternalServiceError,
    ModelUnavailableError,
)
from backend.services.explanation import (  # noqa: E402
    create_evidence_explanation,
    create_match_explanation,
    generate_explanation,
)
from backend.services.history_service import (  # noqa: E402
    get_all_analyses,
    get_analysis_by_id,
)
from backend.services.passage_matcher import find_best_matches  # noqa: E402
from backend.services.passage_selector import (  # noqa: E402
    is_meaningful_passage,
    select_passages,
)
from backend.services.plagiarism_detector import (  # noqa: E402
    classify_match,
    classify_risk,
    detect_candidate_matches,
)
from backend.services.plagiarism_pipeline import (  # noqa: E402
    analyze_documents,
    analyze_web_document,
)
from backend.services.preprocessing import (  # noqa: E402
    normalize_for_tfidf,
    prepare_sentences,
    prepare_text_passages,
)
from backend.services.query_generator import (  # noqa: E402
    generate_queries,
    keyword_query,
)
from backend.services.sbert_model import (  # noqa: E402
    calculate_sbert_similarity,
    generate_embeddings,
    get_model,
    sbert_available,
)
from backend.services.similarity import (  # noqa: E402
    calculate_all_similarities,
    calculate_top_source_matches,
)
from backend.services.source_processor import (  # noqa: E402
    clean_source_text,
    segment_sources,
)
from backend.services.tfidf_model import calculate_tfidf_similarity  # noqa: E402
from backend.services.web_content import (  # noqa: E402
    WebSource,
    retrieve_sources,
    scrape_source,
)
from backend.services.web_search import (  # noqa: E402
    SourceCandidate,
    discover_sources,
    normalize_url,
    search_tavily,
)


class FakeResponse:
    def __init__(self, body, status_code=200):
        self._body = body
        self.status_code = status_code
        self.ok = status_code < 400

    def json(self):
        if isinstance(self._body, Exception):
            raise self._body
        return self._body


class ComprehensiveWebSystemTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        configure_database(f"sqlite:///{TEST_DB_PATH}")
        initialize_database()

    def setUp(self):
        os.environ["TAVILY_API_KEY"] = "mock-tavily-key"
        os.environ["FIRECRAWL_API_KEY"] = "mock-firecrawl-key"
        self.settings = Settings.from_environment()

    # -------------------------------------------------------------
    # 1. Document Extraction & Preprocessing Tests
    # -------------------------------------------------------------
    def test_document_extraction_txt(self):
        txt_file = Path(TEST_DIRECTORY.name) / "sample.txt"
        txt_file.write_text("Sentence one. Sentence two.", encoding="utf-8")
        extracted = extract_text(txt_file)
        self.assertEqual(len(extracted), 1)
        self.assertEqual(extracted[0]["text"], "Sentence one. Sentence two.")

        sentences = process_document(txt_file)
        self.assertEqual(len(sentences), 2)
        self.assertEqual(sentences[0]["text"], "Sentence one.")
        self.assertEqual(sentences[1]["text"], "Sentence two.")

    def test_document_extraction_docx(self):
        docx_file = Path(TEST_DIRECTORY.name) / "sample.docx"
        doc = Document()
        doc.add_paragraph("First paragraph sentence. Second sentence.")
        doc.add_paragraph("Third paragraph sentence.")
        doc.save(docx_file)

        extracted = extract_docx_text(docx_file)
        self.assertEqual(len(extracted), 1)
        self.assertIn("First paragraph sentence", extracted[0]["text"])

        sentences = process_document(docx_file)
        self.assertEqual(len(sentences), 3)

    def test_document_extraction_unsupported_format(self):
        with self.assertRaises(ValueError):
            extract_text(Path(TEST_DIRECTORY.name) / "bad.xyz")

    def test_document_empty_and_normalization(self):
        empty_txt = Path(TEST_DIRECTORY.name) / "empty.txt"
        empty_txt.write_text("   \n\n  ", encoding="utf-8")
        self.assertEqual(extract_txt_text(empty_txt), [])
        self.assertEqual(process_document(empty_txt), [])

        norm = normalize_text("Multiple   spaces\n\nand\tnewlines.")
        self.assertEqual(norm, "Multiple spaces and newlines.")

        tfidf_norm = normalize_for_tfidf("Hello, World! 123 -- Test.")
        self.assertEqual(tfidf_norm, "hello world 123 test")

    # -------------------------------------------------------------
    # 2. Deterministic Passage Selection Tests
    # -------------------------------------------------------------
    def test_passage_selection_filters_headings_and_references(self):
        raw_sentences = [
            {"sentence_id": 1, "page_number": 1, "text": "Introduction"},
            {"sentence_id": 2, "page_number": 1, "text": "Abstract"},
            {"sentence_id": 3, "page_number": 1, "text": "Short"},
            {"sentence_id": 4, "page_number": 1, "text": "References: Smith et al. (2020) Journal of Testing 123-456."},
            {"sentence_id": 5, "page_number": 1, "text": "Deep learning models have demonstrated significant improvements in medical image classification tasks."},
            {"sentence_id": 6, "page_number": 1, "text": "Deep learning models have demonstrated significant improvements in medical image classification tasks."},
            {"sentence_id": 7, "page_number": 1, "text": "123 456 789 000 111 222 333"},
        ]
        selected = select_passages(raw_sentences, maximum=5)
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]["passage_id"], 1)
        self.assertTrue("Deep learning models" in selected[0]["text"])

    def test_is_meaningful_passage_predicates(self):
        self.assertFalse(is_meaningful_passage("Too short."))
        self.assertFalse(is_meaningful_passage("Bibliography"))
        self.assertFalse(is_meaningful_passage("http://example.com/reference-link-to-source-citation"))
        self.assertTrue(is_meaningful_passage("Deep convolutional neural networks are commonly applied to automated feature extraction."))

    # -------------------------------------------------------------
    # 3. Query Generation Tests
    # -------------------------------------------------------------
    def test_query_generation_exact_and_keyword(self):
        passage = {
            "passage_id": 1,
            "text": "Deep learning models have demonstrated significant improvements in medical image classification.",
        }
        queries = generate_queries(passage, maximum_queries=2)
        self.assertEqual(len(queries), 2)
        self.assertIn("medical image classification", queries[0])
        self.assertTrue(any(q.startswith('"') for q in queries))

    # -------------------------------------------------------------
    # 4. URL Normalization & Source Candidate Deduplication
    # -------------------------------------------------------------
    def test_url_normalization(self):
        valid = "https://www.example.com/research/paper?utm_source=twitter&id=42#figure1"
        normalized = normalize_url(valid)
        self.assertEqual(normalized, "https://www.example.com/research/paper?id=42")

        invalid_scheme = "javascript:alert(1)"
        self.assertIsNone(normalize_url(invalid_scheme))

        file_scheme = "file:///etc/passwd"
        self.assertIsNone(normalize_url(file_scheme))

    def test_tavily_search_normalization(self):
        fake_payload = {
            "results": [
                {
                    "title": "Medical Imaging Paper",
                    "url": "https://example.org/med/paper.html?utm_campaign=share#abstract",
                    "content": "Medical image classification overview snippet.",
                    "score": 0.94,
                },
                {
                    "title": "Invalid scheme",
                    "url": "ftp://bad.com/file",
                    "content": "bad",
                    "score": 0.5,
                }
            ]
        }
        with patch("backend.services.web_search.requests.post", return_value=FakeResponse(fake_payload)):
            results = search_tavily("medical imaging", self.settings)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].title, "Medical Imaging Paper")
        self.assertEqual(results[0].url, "https://example.org/med/paper.html")
        self.assertAlmostEqual(results[0].score, 0.94)

    def test_tavily_search_missing_key_and_http_error(self):
        no_key_settings = Settings(
            tavily_api_key=None,
            firecrawl_api_key=None,
            database_url="sqlite:///:memory:",
            max_file_size_mb=10,
            max_search_passages=10,
            max_search_queries=20,
            max_search_results=5,
            max_source_pages=10,
            max_source_text_chars=50000,
            max_source_passages_per_page=200,
            top_k_source_matches=3,
            semantic_weight=0.6,
            lexical_weight=0.4,
            tfidf_direct_threshold=0.85,
            sbert_semantic_threshold=0.75,
            tfidf_lexical_threshold=0.55,
            combined_high_threshold=0.85,
            combined_likely_threshold=0.70,
            combined_possible_threshold=0.50,
            request_timeout=30,
            tavily_search_depth="basic",
        )
        with self.assertRaises(ConfigurationError):
            search_tavily("query", no_key_settings)

        with patch("backend.services.web_search.requests.post", return_value=FakeResponse({"detail": "Rate limit"}, 429)):
            with self.assertRaises(ExternalServiceError) as ctx:
                search_tavily("query", self.settings)
            self.assertEqual(ctx.exception.status_code, 429)

    # -------------------------------------------------------------
    # 5. Firecrawl Scraping & Deduplication Tests
    # -------------------------------------------------------------
    def test_firecrawl_scrape_success_and_failure_handling(self):
        candidate = SourceCandidate(url="https://example.com/article", title="Example Article", search_score=0.91)
        success_payload = {
            "success": True,
            "data": {
                "markdown": "# Article Header\n\nThis is the extracted body text with substantive research information.",
                "metadata": {"title": "Extracted Title", "sourceURL": "https://example.com/article"},
            },
        }
        with patch("backend.services.web_content.requests.post", return_value=FakeResponse(success_payload)):
            source = scrape_source(candidate, self.settings)
        self.assertTrue(source.success)
        self.assertEqual(source.title, "Extracted Title")
        self.assertIn("extracted body text", source.text)

        # Failure per URL does not crash
        with patch("backend.services.web_content.requests.post", return_value=FakeResponse({"detail": "Forbidden"}, 403)):
            failed_source = scrape_source(candidate, self.settings)
        self.assertFalse(failed_source.success)
        self.assertIn("403", failed_source.error)

    def test_retrieve_sources_deduplication(self):
        candidates = [
            SourceCandidate(url="https://example.com/page1", title="Page 1"),
            SourceCandidate(url="https://example.com/page1", title="Page 1 Duplicate"),
            SourceCandidate(url="https://example.com/page2", title="Page 2"),
        ]
        success_payload = {
            "success": True,
            "data": {"markdown": "Meaningful paragraph content for tests.", "metadata": {}},
        }
        with patch("backend.services.web_content.requests.post", return_value=FakeResponse(success_payload)) as mock_post:
            sources = retrieve_sources(candidates, self.settings)
        self.assertEqual(len(sources), 2)
        self.assertEqual(mock_post.call_count, 2)

    # -------------------------------------------------------------
    # 6. Source Processing & Segmentation Tests
    # -------------------------------------------------------------
    def test_clean_source_text_and_segmentation(self):
        raw_markdown = (
            "![banner image](http://example.com/img.png)\n"
            "[Detailed link reference with sufficient words](http://example.com/link)\n"
            "Cookie policy notice: please accept cookies to proceed.\n"
            "# Heading Title\n"
            "This is a valid sentence from the retrieved source with enough length to analyze.\n"
            "Here is another distinct sentence describing machine learning methods."
        )
        cleaned = clean_source_text(raw_markdown, maximum_characters=1000)
        self.assertNotIn("Cookie policy", cleaned)
        self.assertIn("Detailed link reference", cleaned)

        sources = [
            WebSource(
                url="https://example.com/s1",
                title="Source 1",
                text=raw_markdown,
                retrieved_at="2026-01-01T00:00:00Z",
                success=True,
                search_score=0.9,
            ),
            WebSource(
                url="https://example.com/s2",
                title="Source 2",
                text="",
                retrieved_at="2026-01-01T00:00:00Z",
                success=False,
                error="Failed",
            ),
        ]
        passages = segment_sources(sources, maximum_passages_per_page=10, maximum_characters=1000)
        self.assertTrue(len(passages) >= 2)
        self.assertEqual(passages[0]["url"], "https://example.com/s1")

    # -------------------------------------------------------------
    # 7. TF-IDF, SBERT & Evidence Analysis Tests
    # -------------------------------------------------------------
    def test_tfidf_and_sbert_similarity(self):
        s1 = [{"sentence_id": 1, "page_number": 1, "text": "Artificial intelligence transforms industry.", "tfidf_text": "artificial intelligence transforms industry"}]
        s2 = [{"sentence_id": 2, "page_number": 1, "text": "Artificial intelligence transforms industry.", "tfidf_text": "artificial intelligence transforms industry"}]
        s3 = [{"sentence_id": 3, "page_number": 1, "text": "Cooking delicious pasta requires fresh tomatoes.", "tfidf_text": "cooking delicious pasta requires fresh tomatoes"}]

        tfidf_exact = calculate_tfidf_similarity(s1, s2)
        self.assertAlmostEqual(tfidf_exact[0][0], 1.0, places=4)

        tfidf_diff = calculate_tfidf_similarity(s1, s3)
        self.assertAlmostEqual(tfidf_diff[0][0], 0.0, places=4)

        sbert_exact = calculate_sbert_similarity(s1, s2)
        self.assertAlmostEqual(sbert_exact[0][0], 1.0, places=3)

    def test_match_type_and_risk_classification(self):
        # Direct match
        self.assertEqual(classify_match(0.90, 0.90, self.settings), "direct_match")
        # Semantic match (low lexical, high semantic)
        self.assertEqual(classify_match(0.20, 0.85, self.settings), "semantic_match")
        # Lexical overlap (high lexical, lower semantic)
        self.assertEqual(classify_match(0.65, 0.50, self.settings), "lexical_overlap")
        # Low similarity
        self.assertEqual(classify_match(0.10, 0.20, self.settings), "low_similarity")

        # Risk level classification
        self.assertEqual(classify_risk(0.90, self.settings), "highly_suspicious")
        self.assertEqual(classify_risk(0.75, self.settings), "likely_plagiarism")
        self.assertEqual(classify_risk(0.55, self.settings), "possibly_similar")
        self.assertEqual(classify_risk(0.30, self.settings), "no_strong_evidence")

    def test_evidence_analysis_and_explanation(self):
        student = {"text": "Machine learning identifies disease patterns in clinical radiography.", "page": 1, "position": 1}
        source = {"title": "Radiography ML Study", "url": "https://journal.example/ml-radio", "text": "Machine learning identifies disease patterns in radiography.", "position": 3, "search_score": 0.88}

        evidence = analyze_evidence(student, source, tfidf_score=0.80, sbert_score=0.92, settings=self.settings)
        self.assertEqual(evidence["match_type"], "semantic_match")
        self.assertEqual(evidence["risk"], "highly_suspicious")
        self.assertIn("TF-IDF similarity is 80.0%", evidence["explanation"])
        self.assertIn("not a definitive academic or legal judgment", evidence["explanation"])

    # -------------------------------------------------------------
    # 8. Database Persistence & History Service Tests
    # -------------------------------------------------------------
    def test_save_and_retrieve_web_analysis(self):
        report = {
            "filename": "thesis_chapter.pdf",
            "status": "completed",
            "summary": {
                "total_passages": 25,
                "searched_passages": 8,
                "sources_found": 3,
                "sources_retrieved": 3,
                "suspicious_matches": 1,
                "overall_risk": "likely_plagiarism",
            },
            "matches": [
                {
                    "submitted_passage": "Student sentence text.",
                    "page": 2,
                    "position": 5,
                    "source": {"title": "Source Article", "url": "https://example.com/source"},
                    "matched_passage": "Source sentence text.",
                    "source_position": 1,
                    "similarity": {"tfidf": 0.72, "sbert": 0.88, "combined": 0.816},
                    "match_type": "semantic_match",
                    "risk": "likely_plagiarism",
                    "explanation": "Evidence explanation text.",
                }
            ],
            "sources": [
                {
                    "title": "Source Article",
                    "url": "https://example.com/source",
                    "retrieval_status": "success",
                    "error": None,
                    "search_score": 0.95,
                }
            ],
            "warnings": [],
        }
        analysis_id = save_web_analysis_to_database(report)
        self.assertIsInstance(analysis_id, int)

        stored = get_analysis_by_id(analysis_id)
        self.assertIsNotNone(stored)
        self.assertEqual(stored["filename"], "thesis_chapter.pdf")
        self.assertEqual(stored["summary"]["overall_risk"], "likely_plagiarism")
        self.assertEqual(len(stored["matches"]), 1)
        self.assertEqual(stored["matches"][0]["source"]["url"], "https://example.com/source")
        self.assertEqual(stored["matches"][0]["similarity"]["combined"], 0.816)
        self.assertEqual(len(stored["sources"]), 1)

        all_records = get_all_analyses()
        self.assertTrue(any(item["id"] == analysis_id for item in all_records))

    def test_save_and_retrieve_legacy_two_document_analysis(self):
        legacy_result = {
            "source_document": "docA.pdf",
            "submitted_document": "docB.pdf",
            "source_sentence_count": 2,
            "submitted_sentence_count": 2,
            "matched_sentence_count": 1,
            "matches": [
                {
                    "source_sentence_id": 1,
                    "submitted_sentence_id": 1,
                    "source_page": 1,
                    "submitted_page": 1,
                    "source_text": "Sample source.",
                    "submitted_text": "Sample submitted.",
                    "tfidf_score": 0.9,
                    "sbert_score": 0.95,
                    "match_type": "direct_match",
                    "explanation": "Direct match explanation.",
                }
            ]
        }
        legacy_id = save_analysis_to_database(legacy_result)
        stored_legacy = get_analysis_by_id(legacy_id)
        self.assertIsNotNone(stored_legacy)
        self.assertEqual(stored_legacy["source_document"], "docA.pdf")
        self.assertEqual(len(stored_legacy["matches"]), 1)

    # -------------------------------------------------------------
    # 9. Complete Mocked Pipeline End-to-End Test
    # -------------------------------------------------------------
    def test_full_mocked_web_pipeline(self):
        doc_path = Path(TEST_DIRECTORY.name) / "e2e_student.txt"
        doc_path.write_text(
            "Natural language processing helps computers understand human language.\n"
            "Semantic similarity measures how closely two pieces of text are related.",
            encoding="utf-8",
        )
        mock_candidates = [
            SourceCandidate(
                url="https://nlp-hub.org/intro",
                title="Introduction to NLP",
                search_score=0.95,
                matched_queries=['"Natural language processing helps computers understand human language."'],
                matched_passage_ids=[1],
            )
        ]
        mock_web_source = WebSource(
            url="https://nlp-hub.org/intro",
            title="Introduction to NLP",
            text="Natural language processing enables computers to comprehend human languages effectively.",
            retrieved_at="2026-01-01T00:00:00Z",
            success=True,
            search_score=0.95,
        )
        fake_sbert_matrix = np.array([[0.92, 0.40]])

        with patch("backend.services.plagiarism_pipeline.discover_sources", return_value=mock_candidates), \
             patch("backend.services.plagiarism_pipeline.retrieve_sources", return_value=[mock_web_source]), \
             patch("backend.services.similarity.calculate_sbert_similarity", return_value=fake_sbert_matrix):
            report = analyze_web_document(doc_path, "e2e_student.txt", self.settings)

        self.assertEqual(report["status"], "completed")
        self.assertEqual(report["summary"]["sources_found"], 1)
        self.assertEqual(report["summary"]["sources_retrieved"], 1)
        self.assertEqual(len(report["sources"]), 1)
        self.assertTrue(len(report["matches"]) >= 1)

        analysis_id = save_web_analysis_to_database(report)
        persisted = get_analysis_by_id(analysis_id)
        self.assertEqual(persisted["filename"], "e2e_student.txt")

    # -------------------------------------------------------------
    # 10. Flask API Endpoints & Validation Tests
    # -------------------------------------------------------------
    def test_flask_home_and_health(self):
        client = app.test_client()
        res_home = client.get("/")
        self.assertEqual(res_home.status_code, 200)
        self.assertEqual(res_home.json["workflow"], "one document -> public web sources -> evidence report")

        res_health = client.get("/api/health")
        self.assertEqual(res_health.status_code, 200)
        self.assertEqual(res_health.json["status"], "healthy")
        self.assertTrue(res_health.json["tavily_configured"])
        self.assertTrue(res_health.json["firecrawl_configured"])

    def test_flask_analyze_validations(self):
        client = app.test_client()

        # Missing 'document'
        res_empty = client.post("/api/analyze", data={})
        self.assertEqual(res_empty.status_code, 400)

        # Bad extension
        res_bad_ext = client.post(
            "/api/analyze",
            data={"document": (io.BytesIO(b"data"), "image.png")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_bad_ext.status_code, 400)

        # Empty file
        res_empty_file = client.post(
            "/api/analyze",
            data={"document": (io.BytesIO(b""), "empty.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(res_empty_file.status_code, 400)

    def test_flask_analyze_successful_mocked_upload(self):
        client = app.test_client()
        mock_report = {
            "filename": "uploaded.txt",
            "status": "completed",
            "summary": {
                "total_passages": 2,
                "searched_passages": 1,
                "sources_found": 1,
                "sources_retrieved": 1,
                "suspicious_matches": 0,
                "overall_risk": "no_strong_evidence",
            },
            "matches": [],
            "sources": [{"title": "Web Source", "url": "https://example.com", "retrieval_status": "success", "error": None, "search_score": 0.8}],
            "warnings": [],
        }
        with patch("backend.app.analyze_web_document", return_value=mock_report):
            response = client.post(
                "/api/analyze",
                data={"document": (io.BytesIO(b"Valid student paper content for automated plagiarism checking."), "uploaded.txt")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("analysis_id", response.json)
        self.assertEqual(response.json["filename"], "uploaded.txt")

    def test_flask_history_endpoints(self):
        client = app.test_client()
        res_list = client.get("/api/history")
        self.assertEqual(res_list.status_code, 200)
        self.assertIn("analyses", res_list.json)

        res_404 = client.get("/api/history/999999")
        self.assertEqual(res_404.status_code, 404)

        # Test delete single
        if res_list.json["analyses"]:
            target_id = res_list.json["analyses"][0]["id"]
            res_del_single = client.delete(f"/api/history/{target_id}")
            self.assertEqual(res_del_single.status_code, 200)

        # Test clear all history
        res_clear = client.delete("/api/history")
        self.assertEqual(res_clear.status_code, 200)
        res_after = client.get("/api/history")
        self.assertEqual(len(res_after.json["analyses"]), 0)


if __name__ == "__main__":
    unittest.main()
